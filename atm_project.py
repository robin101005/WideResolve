"""
Mini Project: ATM System with Excel Sheet Integration
------------------------------------------------------
This program stores account details (Name, Contact, Balance) in an
Excel file (atm_data.xlsx) using the openpyxl library.

Every time you deposit or withdraw money, the balance in the Excel
sheet is automatically updated. If you try to withdraw more money
than your balance, the program will notify you and stop the
transaction.
"""

import openpyxl
from openpyxl import Workbook
import os

FILE_NAME = "atm_data.xlsx"
SHEET_NAME = "Accounts"
HEADERS = ["Account Number", "Name", "Contact", "Balance"]


# ---------- Excel Helper Functions ----------

def setup_excel():
    """Create the Excel file with headers if it doesn't already exist."""
    if not os.path.exists(FILE_NAME):
        wb = Workbook()
        sheet = wb.active
        sheet.title = SHEET_NAME
        sheet.append(HEADERS)
        wb.save(FILE_NAME)


def load_workbook():
    return openpyxl.load_workbook(FILE_NAME)


def find_account_row(sheet, account_number):
    """Return the row number where this account number is stored, or None."""
    for row in range(2, sheet.max_row + 1):
        if str(sheet.cell(row=row, column=1).value) == str(account_number):
            return row
    return None


# ---------- ATM Operations ----------

def create_account():
    account_number = input("Enter new account number: ").strip()
    wb = load_workbook()
    sheet = wb[SHEET_NAME]

    if find_account_row(sheet, account_number):
        print("An account with this number already exists.\n")
        return

    name = input("Enter account holder name: ").strip()
    contact = input("Enter contact number: ").strip()

    try:
        opening_balance = float(input("Enter opening deposit amount: "))
        if opening_balance < 0:
            print("Opening balance cannot be negative.\n")
            return
    except ValueError:
        print("Invalid amount entered.\n")
        return

    sheet.append([account_number, name, contact, opening_balance])
    wb.save(FILE_NAME)
    print(f"Account created successfully for {name}. Balance: {opening_balance}\n")


def deposit_money():
    account_number = input("Enter account number: ").strip()
    wb = load_workbook()
    sheet = wb[SHEET_NAME]
    row = find_account_row(sheet, account_number)

    if not row:
        print("Account not found.\n")
        return

    try:
        amount = float(input("Enter amount to deposit: "))
        if amount <= 0:
            print("Deposit amount must be positive.\n")
            return
    except ValueError:
        print("Invalid amount entered.\n")
        return

    current_balance = sheet.cell(row=row, column=4).value
    new_balance = current_balance + amount
    sheet.cell(row=row, column=4).value = new_balance
    wb.save(FILE_NAME)

    print(f"Deposit successful! New balance: {new_balance}\n")


def withdraw_money():
    account_number = input("Enter account number: ").strip()
    wb = load_workbook()
    sheet = wb[SHEET_NAME]
    row = find_account_row(sheet, account_number)

    if not row:
        print("Account not found.\n")
        return

    try:
        amount = float(input("Enter amount to withdraw: "))
        if amount <= 0:
            print("Withdrawal amount must be positive.\n")
            return
    except ValueError:
        print("Invalid amount entered.\n")
        return

    current_balance = sheet.cell(row=row, column=4).value

    # This is the key check the teacher asked for:
    # notify the user if withdrawal amount is higher than balance
    if amount > current_balance:
        print(f"Insufficient balance! Your current balance is only {current_balance}.\n")
        return

    new_balance = current_balance - amount
    sheet.cell(row=row, column=4).value = new_balance
    wb.save(FILE_NAME)

    print(f"Withdrawal successful! New balance: {new_balance}\n")


def check_balance():
    account_number = input("Enter account number: ").strip()
    wb = load_workbook()
    sheet = wb[SHEET_NAME]
    row = find_account_row(sheet, account_number)

    if not row:
        print("Account not found.\n")
        return

    name = sheet.cell(row=row, column=2).value
    balance = sheet.cell(row=row, column=4).value
    print(f"Account Holder: {name}\nCurrent Balance: {balance}\n")


def view_all_accounts():
    wb = load_workbook()
    sheet = wb[SHEET_NAME]

    if sheet.max_row < 2:
        print("No accounts found.\n")
        return

    print(f"{'Acc No':<12}{'Name':<20}{'Contact':<15}{'Balance':<10}")
    print("-" * 57)
    for row in range(2, sheet.max_row + 1):
        acc = sheet.cell(row=row, column=1).value
        name = sheet.cell(row=row, column=2).value
        contact = sheet.cell(row=row, column=3).value
        balance = sheet.cell(row=row, column=4).value
        print(f"{str(acc):<12}{str(name):<20}{str(contact):<15}{str(balance):<10}")
    print()


# ---------- Main Menu ----------

def main():
    setup_excel()

    while True:
        print("===== ATM MENU =====")
        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Check Balance")
        print("5. View All Accounts")
        print("6. Exit")

        choice = input("Enter your choice (1-6): ").strip()
        print()

        if choice == "1":
            create_account()
        elif choice == "2":
            deposit_money()
        elif choice == "3":
            withdraw_money()
        elif choice == "4":
            check_balance()
        elif choice == "5":
            view_all_accounts()
        elif choice == "6":
            print("Thank you for using the ATM. Goodbye!")
            break
        else:
            print("Invalid choice. Please try again.\n")


if __name__ == "__main__":
    main()